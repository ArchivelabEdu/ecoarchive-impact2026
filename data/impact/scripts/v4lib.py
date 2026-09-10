"""v0.1/v0.3 서식과 동일한 XML 요소를 생성하는 헬퍼 (python-docx + lxml)"""
import copy, re
from lxml import etree
from docx.oxml.ns import qn
from docx.shared import Pt, Emu

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = ('xmlns:w="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"' % W)

DARK, MED, GRAY, TEXT = '2E6E3F', '4C9A5A', '5F6368', '2B2B2B'
TABLE_BORDER, TABLE_ALT, BOX_BG, QUOTE_BG, QUOTE_TEXT = 'E4EBE5', 'F1F7F2', 'EAF3EC', 'F6FAF6', '3A5A43'
CONTENT_W = 9026

def esc(s):
    return (str(s).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;'))

def mk(xml):
    return etree.fromstring(f'<w:root {NS}>{xml}</w:root>')[0]

def rpr(bold=False, italic=False, color=None, size=None):
    s = ''
    if bold: s += '<w:b/><w:bCs/>'
    if italic: s += '<w:i/><w:iCs/>'
    if color: s += f'<w:color w:val="{color}"/>'
    if size: s += f'<w:sz w:val="{size}"/><w:szCs w:val="{size}"/>'
    return f'<w:rPr>{s}</w:rPr>' if s else ''

def run(text, **k):
    return f'<w:r>{rpr(**k)}<w:t xml:space="preserve">{esc(text)}</w:t></w:r>'

def para(runs_xml, ppr=''):
    return mk(f'<w:p>{("<w:pPr>" + ppr + "</w:pPr>") if ppr else ""}{runs_xml}</w:p>')

# ---- 문단 요소 ----
def body(text_or_runs):
    r = run(text_or_runs) if isinstance(text_or_runs, str) else ''.join(text_or_runs)
    return para(r, '<w:spacing w:after="140" w:line="300" w:lineRule="auto"/>')

def chapter(title):
    m = re.match(r'^(\d+\.\s+|부록\.\s+)(.*)$', title)
    runs = (run(m.group(1), bold=True, size=26, color=MED) + run(m.group(2), bold=True, size=26, color=DARK)) if m \
        else run(title, bold=True, size=26, color=DARK)
    return para(runs, f'<w:keepNext/><w:pBdr><w:bottom w:val="single" w:sz="12" w:space="6" w:color="{MED}"/></w:pBdr>'
                      '<w:spacing w:before="320" w:after="160" w:line="300" w:lineRule="auto"/>')

def h2(title):
    return para(run(title, bold=True, size=22), '<w:keepNext/><w:spacing w:before="240" w:after="100" w:line="300" w:lineRule="auto"/>')

def h3(title):
    return para(run(title, bold=True, size=20), '<w:keepNext/><w:spacing w:before="200" w:after="90" w:line="300" w:lineRule="auto"/>')

def caption(text):
    return para(run(text, italic=True, color=GRAY, size=16), '<w:spacing w:after="160"/><w:jc w:val="center"/>')

def footnote(text):
    return para(run(text, italic=True, color=GRAY, size=16), '<w:spacing w:before="80" w:after="140"/>')

def placeholder(text):
    return para(run(text, italic=True, color=GRAY, size=19), '<w:spacing w:after="130" w:line="300" w:lineRule="auto"/>')

def spacer(after=160):
    return para('', f'<w:spacing w:after="{after}"/>')

def page_break():
    return mk('<w:p><w:r><w:br w:type="page"/></w:r></w:p>')

def bullet(runs, num_id='1'):
    r = run(runs) if isinstance(runs, str) else ''.join(runs)
    return para(r, f'<w:pStyle w:val="a4"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
                   '<w:spacing w:after="80" w:line="288" w:lineRule="auto"/>')

def bt(label, rest):
    """bold label + text runs"""
    return [run(label, bold=True), run(rest)]

def quote_para(text):
    return para(run(text, italic=True, size=19, color=QUOTE_TEXT), '<w:spacing w:line="288" w:lineRule="auto"/>')

# ---- 표 ----
def _tblpr(borders=True):
    if borders:
        b = ''.join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="{TABLE_BORDER}"/>'
                    for s in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV'])
    else:
        b = '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/><w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
    return (f'<w:tblPr><w:tblW w:w="{CONTENT_W}" w:type="dxa"/><w:tblBorders>{b}</w:tblBorders>'
            '<w:tblCellMar><w:left w:w="10" w:type="dxa"/><w:right w:w="10" w:type="dxa"/></w:tblCellMar>'
            '<w:tblLook w:val="0000" w:firstRow="0" w:lastRow="0" w:firstColumn="0" w:lastColumn="0" w:noHBand="0" w:noVBand="0"/></w:tblPr>')

def _grid(widths):
    return '<w:tblGrid>' + ''.join(f'<w:gridCol w:w="{w}"/>' for w in widths) + '</w:tblGrid>'

def _cell(content_xml, width, fill, align='left', bold=False, color=None, size=18, mar=(60, 90, 60, 90), extra_tcpr=''):
    jc = '' if align == 'left' else f'<w:jc w:val="{align}"/>'
    if isinstance(content_xml, str):  # plain text (\n -> multiple paragraphs)
        paras = ''.join(f'<w:p><w:pPr><w:spacing w:line="264" w:lineRule="auto"/>{jc}</w:pPr>'
                        f'{run(ln, bold=bold, color=color, size=size)}</w:p>' for ln in str(content_xml).split('\n'))
    else:
        paras = content_xml
    t, l, b, r = mar
    return (f'<w:tc><w:tcPr><w:tcW w:w="{width}" w:type="dxa"/>{extra_tcpr}'
            f'<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>'
            f'<w:tcMar><w:top w:w="{t}" w:type="dxa"/><w:left w:w="{l}" w:type="dxa"/><w:bottom w:w="{b}" w:type="dxa"/><w:right w:w="{r}" w:type="dxa"/></w:tcMar>'
            f'<w:vAlign w:val="center"/></w:tcPr>{paras}</w:tc>')

def table(widths, headers, rows, aligns=None, bold_last=False, header_align_first='left'):
    """v0.1 표 서식: 헤더 진초록/흰글씨, 교차행 연두, 얇은 그리드. aligns: 'L'/'C' per column"""
    aligns = aligns or (['L'] + ['C'] * (len(widths) - 1))
    al = lambda a: 'left' if a == 'L' else 'center'
    xml = '<w:tbl>' + _tblpr() + _grid(widths)
    xml += '<w:tr><w:trPr><w:cantSplit/><w:tblHeader/></w:trPr>'
    for i, h in enumerate(headers):
        xml += _cell(h, widths[i], DARK, header_align_first if i == 0 else 'center', color='FFFFFF')
    xml += '</w:tr>'
    for ri, r in enumerate(rows):
        fill = TABLE_ALT if ri % 2 == 1 else 'FFFFFF'
        bold = bold_last and ri == len(rows) - 1
        xml += '<w:tr><w:trPr><w:cantSplit/></w:trPr>'
        for ci, c in enumerate(r):
            xml += _cell(c, widths[ci], fill, al(aligns[ci]), bold=bold)
        xml += '</w:tr>'
    return mk(xml + '</w:tbl>')

def box(title, paragraphs):
    """인사이트/시사점 박스"""
    inner = f'<w:p><w:pPr><w:spacing w:after="80"/></w:pPr>{run(title, bold=True, color=DARK)}</w:p>'
    for p in paragraphs:
        r = run(p, size=19) if isinstance(p, str) else ''.join(p)
        inner += f'<w:p><w:pPr><w:spacing w:line="288" w:lineRule="auto"/></w:pPr>{r}</w:p>'
    cell = (f'<w:tc><w:tcPr><w:tcW w:w="{CONTENT_W}" w:type="dxa"/>'
            f'<w:tcBorders><w:left w:val="single" w:sz="24" w:space="0" w:color="{MED}"/></w:tcBorders>'
            f'<w:shd w:val="clear" w:color="auto" w:fill="{BOX_BG}"/>'
            '<w:tcMar><w:top w:w="140" w:type="dxa"/><w:left w:w="200" w:type="dxa"/><w:bottom w:w="140" w:type="dxa"/><w:right w:w="200" w:type="dxa"/></w:tcMar>'
            f'</w:tcPr>{inner}</w:tc>')
    return mk('<w:tbl>' + _tblpr(False) + _grid([CONTENT_W]) + f'<w:tr><w:trPr><w:cantSplit/></w:trPr>{cell}</w:tr></w:tbl>')

def quote_box(text):
    cell = (f'<w:tc><w:tcPr><w:tcW w:w="{CONTENT_W}" w:type="dxa"/>'
            f'<w:tcBorders><w:left w:val="single" w:sz="20" w:space="0" w:color="{MED}"/></w:tcBorders>'
            f'<w:shd w:val="clear" w:color="auto" w:fill="{QUOTE_BG}"/>'
            '<w:tcMar><w:top w:w="100" w:type="dxa"/><w:left w:w="200" w:type="dxa"/><w:bottom w:w="100" w:type="dxa"/><w:right w:w="160" w:type="dxa"/></w:tcMar>'
            f'</w:tcPr><w:p><w:pPr><w:spacing w:line="288" w:lineRule="auto"/></w:pPr>{run(text, italic=True, size=19, color=QUOTE_TEXT)}</w:p></w:tc>')
    return mk('<w:tbl>' + _tblpr(False) + _grid([CONTENT_W]) + f'<w:tr><w:trPr><w:cantSplit/></w:trPr>{cell}</w:tr></w:tbl>')

# ---- 이미지 ----
def image_para(doc, path, width_px, height_px=None):
    """문서에 이미지를 추가하고 중앙정렬 문단 요소를 반환 (원본 이미지 문단 서식과 동일)"""
    p = doc.add_paragraph()
    r = p.add_run()
    from PIL import Image
    if height_px is None:
        im = Image.open(path); height_px = int(width_px * im.height / im.width)
    r.add_picture(path, width=Emu(width_px * 9525), height=Emu(height_px * 9525))
    el = p._p
    el.getparent().remove(el)
    el.insert(0, mk('<w:pPr><w:spacing w:before="120" w:after="40"/><w:jc w:val="center"/></w:pPr>'))
    return el

def drawing_run_from(el, width_px=None, height_px=None):
    """기존 요소에서 <w:drawing>을 포함한 run을 deepcopy (필요시 크기 조정)"""
    for r in el.iter(qn('w:r')):
        if r.find(qn('w:drawing')) is not None:
            r2 = copy.deepcopy(r)
            if width_px:
                cx, cy = width_px * 9525, height_px * 9525
                for tag in ['{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent',
                            '{http://schemas.openxmlformats.org/drawingml/2006/main}ext']:
                    for e in r2.iter(tag):
                        e.set('cx', str(cx)); e.set('cy', str(cy))
            return r2
    return None

def image_para_from_run(run_el):
    p = mk('<w:p><w:pPr><w:spacing w:before="120" w:after="40"/><w:jc w:val="center"/></w:pPr></w:p>')
    p.append(run_el)
    return p

def set_para_text(p_el, new_runs_xml):
    """문단의 기존 run들을 제거하고 새 run(xml 문자열)으로 교체 (pPr 유지)"""
    for ch in list(p_el):
        if ch.tag != qn('w:pPr'):
            p_el.remove(ch)
    for r in etree.fromstring(f'<w:root {NS}>{new_runs_xml}</w:root>'):
        p_el.append(r)

def ptext(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))
